import argparse, json, os, sys
import hashlib
import re
from pathlib import Path
from datetime import datetime

SCHEMA_VERSION = 1
CHAPTER_NUMBER_RE = re.compile(r"第(\d+)")
PROJECT_DIRS = ("设定集", "大纲", "正文", "追踪", "审查报告", "拆文库")
DATA_ARTIFACT_NAMES = (
    "fulfillment_result.json",
    "disambiguation_result.json",
    "extraction_result.json",
)


def chapter_number(path):
    """Return a numeric chapter number for a conventional chapter filename."""
    match = CHAPTER_NUMBER_RE.search(path.name)
    return int(match.group(1)) if match else 10**9


def chapter_files(root):
    """Return chapter Markdown files in numeric, then lexical order."""
    directory = root / "正文"
    if not directory.exists():
        return []
    return sorted(directory.glob("第*.md"), key=lambda p: (chapter_number(p), p.name))

def find_project(start):
    """只读查找项目根目录，找不到返回 None，不创建任何文件"""
    c = Path(start).resolve()
    if not c.exists() or not c.is_dir():
        return None
    state_file = c / ".novel" / "state.json"
    if state_file.exists():
        return c
    for child in sorted(c.iterdir()):
        if child.is_dir():
            child_state = child / ".novel" / "state.json"
            if child_state.exists():
                return child
    return None

def migrate_state(state):
    """迁移旧版 state.json 到当前 schema。

    历史版本：
      - 无 schema_version 字段：早期版本，可能用 'project_info' 而非 'project'
      - schema_version=1：当前版本，统一用 'project'

    返回 (migrated_state, changed) 元组。changed=True 表示发生了迁移需要持久化。
    幂等，可重复调用。
    """
    if not isinstance(state, dict):
        return state, False
    # 已是当前版本
    if state.get("schema_version") == SCHEMA_VERSION:
        return state, False
    # 旧版迁移：project_info -> project
    if "project_info" in state and "project" not in state:
        state["project"] = state.pop("project_info")
    state["schema_version"] = SCHEMA_VERSION
    return state, True

def init_project(start, title="", author="", genre=""):
    """显式创建新项目，返回项目根"""
    c = Path(start).resolve()
    novel_dir = c / ".novel"
    novel_dir.mkdir(parents=True, exist_ok=True)
    state_file = novel_dir / "state.json"
    if state_file.exists():
        # 幂等补齐运行时目录，但不覆盖用户已有状态或 idea bank。
        for dirname in PROJECT_DIRS:
            (c / dirname).mkdir(parents=True, exist_ok=True)
        (novel_dir / "tmp").mkdir(parents=True, exist_ok=True)
        if not (novel_dir / "idea_bank.json").exists():
            (novel_dir / "idea_bank.json").write_text(
                json.dumps({"ideas": []}, ensure_ascii=False, indent=2), "utf-8"
            )
        if not (novel_dir / "active-book").exists():
            (novel_dir / "active-book").write_text(c.name, encoding="utf-8")
        return c

    for dirname in PROJECT_DIRS:
        (c / dirname).mkdir(parents=True, exist_ok=True)
    (novel_dir / "tmp").mkdir(parents=True, exist_ok=True)

    init_state = {
        "schema_version": SCHEMA_VERSION,
        "project": {
            "book_name": title,
            "author": author,
            "genre": genre,
            "target_words": 0,
            "target_platform": "fanqie",
        },
        "progress": {
            "current_chapter": 0,
            "current_volume": 1,
            "writing_started": False,
        },
        "versions": {"baseline_version": 0, "last_review_chapter": 0},
    }
    state_file.write_text(json.dumps(init_state, ensure_ascii=False, indent=2), "utf-8")
    (novel_dir / "idea_bank.json").write_text(
        json.dumps({"ideas": []}, ensure_ascii=False, indent=2), "utf-8"
    )
    (novel_dir / "active-book").write_text(c.name, encoding="utf-8")
    return c

def load_state(r):
    f = r / ".novel" / "state.json"
    if f.exists():
        try:
            state = json.loads(f.read_text("utf-8"))
            # 自动迁移旧版 schema（含 project_info 兼容），持久化迁移结果
            migrated, changed = migrate_state(state)
            if changed:
                save_state(r, migrated)
            return migrated
        except Exception:
            return {}
    return {}

def save_state(r, s):
    f = r / ".novel" / "state.json"
    f.parent.mkdir(parents=True, exist_ok=True)
    f.write_text(json.dumps(s, ensure_ascii=False, indent=2), "utf-8")


def markdown_to_plain_text(text):
    """Remove common Markdown presentation markers while preserving prose."""
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = normalized.split("\n")
    cleaned = []
    in_frontmatter = False
    in_fence = False
    for index, line in enumerate(lines):
        stripped = line.strip()
        if index == 0 and stripped == "---":
            in_frontmatter = True
            continue
        if in_frontmatter:
            if stripped == "---":
                in_frontmatter = False
            continue
        if stripped.startswith("```") or stripped.startswith("~~~"):
            in_fence = not in_fence
            continue
        if in_fence:
            cleaned.append(line)
            continue
        if re.match(r"^\s{0,3}(?:[-*_]\s*){3,}$", line):
            continue
        line = re.sub(r"^\s{0,3}#{1,6}\s*", "", line)
        line = re.sub(r"^\s*>\s?", "", line)
        line = re.sub(r"^\s*(?:[-+*]|\d+[.)])\s+", "", line)
        line = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", line)
        line = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", line)
        line = re.sub(r"<[^>\n]+>", "", line)
        line = re.sub(r"(\*\*|__|~~|`)", "", line)
        line = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"\1", line)
        line = re.sub(r"(?<!_)_([^_\n]+)_(?!_)", r"\1", line)
        cleaned.append(line.rstrip())
    return "\n".join(cleaned).strip()


def load_review_results(value):
    """Load reviewer JSON from a file path or inline JSON.

    Returns ``(payload, source, error)``. Invalid or empty results return a
    falsey payload and an explanatory error so callers cannot mark a review as
    verified merely because a path string was supplied.
    """
    if not value:
        return None, None, None

    source = "inline"
    raw = value
    try:
        candidate = Path(value).expanduser()
        if candidate.is_file():
            source = "file:" + str(candidate)
            raw = candidate.read_text(encoding="utf-8")
    except (OSError, ValueError) as exc:
        return None, source, "cannot read review results: " + str(exc)

    try:
        payload = json.loads(raw)
        if isinstance(payload, str):
            payload = json.loads(payload)
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return None, source, "invalid review results JSON: " + str(exc)

    if not isinstance(payload, (dict, list)) or not payload:
        return None, source, "review results must be a non-empty JSON object or array"
    return payload, source, None


def review_results_verified(payload):
    """Return whether the reviewer explicitly supplied usable results."""
    if not payload:
        return False
    if isinstance(payload, dict):
        status = str(payload.get("status", "")).lower()
        if status in {"incomplete", "failed", "unverified"}:
            return False
    return True


def collect_data_artifacts(r, source=None):
    """Read the three optional data-agent artifacts without inventing facts."""
    base = Path(source).expanduser() if source else r / ".novel" / "tmp"
    artifacts = {}
    errors = []

    if base.is_file():
        try:
            bundle = json.loads(base.read_text(encoding="utf-8"))
            if not isinstance(bundle, dict):
                errors.append("artifact bundle must be a JSON object")
            else:
                for name in DATA_ARTIFACT_NAMES:
                    if name in bundle:
                        artifacts[name] = bundle[name]
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            errors.append(str(base) + ": " + str(exc))
    elif base.is_dir():
        for name in DATA_ARTIFACT_NAMES:
            path = base / name
            if not path.exists():
                continue
            try:
                artifacts[name] = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, UnicodeError, json.JSONDecodeError) as exc:
                errors.append(name + ": " + str(exc))
    elif source:
        errors.append("data-artifacts path not found: " + str(base))

    if len(artifacts) == len(DATA_ARTIFACT_NAMES) and not errors:
        status = "complete"
    elif artifacts or errors:
        status = "partial"
    else:
        status = "missing"
    return status, artifacts, errors


def project_chapter_summary(r, chapter, commit):
    """Refresh the deterministic summary projection for one committed chapter."""
    source_path = commit.get("source_path")
    if not source_path:
        raise ValueError("commit has no source_path")
    root = r.resolve()
    source = (r / source_path).resolve()
    try:
        source.relative_to(root)
    except ValueError as exc:
        raise ValueError("commit source_path escapes project root") from exc
    if not source.is_file():
        raise FileNotFoundError(str(source))

    body = markdown_to_plain_text(source.read_text(encoding="utf-8"))
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", body) if part.strip()]
    summary = paragraphs[0] if paragraphs else ""
    if len(summary) > 280:
        summary = summary[:277].rstrip() + "..."

    summary_dir = r / ".story-system" / "summaries"
    summary_dir.mkdir(parents=True, exist_ok=True)
    summary_file = summary_dir / f"ch{int(chapter):04d}.md"
    summary_file.write_text(
        "# 第" + str(chapter) + "章摘要\n\n"
        + "来源: " + source_path + "\n"
        + "正文 SHA-256: " + str(commit.get("content_sha256") or "") + "\n\n"
        + summary + "\n",
        encoding="utf-8",
    )
    return summary_file

def cmd_init(a):
    r = init_project(Path(a.project_root), a.title or "", a.author or "", a.genre or "")
    print(str(r))

def cmd_where(a):
    r = find_project(Path(a.project_root))
    if r:
        print(str(r))
    else:
        print("no project found", file=sys.stderr)
        sys.exit(1)

def cmd_status(a):
    r = find_project(Path(a.project_root))
    if not r:
        print("no project")
        return
    s = load_state(r)
    if not s:
        print("no state")
        return
    # load_state 已自动迁移 'project_info' -> 'project'，这里直接用 'project'
    pi = s.get("project", {})
    pr = s.get("progress", {})
    payload = {
        "project_root": str(r),
        "book": pi.get("book_name", pi.get("title", "?")),
        "genre": pi.get("genre", "?"),
        "writing_started": pr.get("writing_started", False),
        "current_chapter": pr.get("current_chapter", 0),
        "current_volume": pr.get("current_volume", 1),
        "phase": "writing" if pr.get("writing_started", False) else "prewriting",
    }
    if getattr(a, "format", "summary") == "json":
        print(json.dumps(payload, ensure_ascii=False))
        return
    print("book:", pi.get("book_name", pi.get("title", "?")))
    print("genre:", pi.get("genre", "?"))
    print("writing_started:", pr.get("writing_started", False))
    print("current_chapter:", pr.get("current_chapter", 0))
    print("phase:", payload["phase"])

def cmd_doctor(a):
    output_json = getattr(a, "format", "text") == "json"
    r = find_project(Path(a.project_root))
    if not r:
        if output_json:
            print(json.dumps({"issues": ["project not found"], "ok": False, "chapters": 0}, ensure_ascii=False))
        else:
            print("no project")
        return 1
    issues = []
    try:
        state_file = r / ".novel" / "state.json"
        if state_file.exists():
            try:
                state = json.loads(state_file.read_text("utf-8"))
                if state.get("schema_version") != SCHEMA_VERSION:
                    issues.append("state schema version is not current")
            except (OSError, json.JSONDecodeError) as exc:
                issues.append("invalid: .novel/state.json (" + str(exc) + ")")

        checks = [
            (r / ".novel" / "state.json", "state.json"),
            (r / "设定集" / "主角卡.md", "设定集/主角卡.md"),
            (r / "设定集" / "世界观.md", "设定集/世界观.md"),
        ]
        started = bool(load_state(r).get("progress", {}).get("writing_started", False))
        checks.append(
            (r / "大纲" / ("总纲.md" if started else "总纲草案.md"),
             "大纲/总纲.md" if started else "大纲/总纲草案.md")
        )
        for path_obj, name in checks:
            if not path_obj.exists():
                issues.append("missing: " + name)
        chs = chapter_files(r)
        if chs and not output_json:
            print("chapters:", len(chs))
        if started or chs:
            tracks = ["追踪/上下文.md", "追踪/伏笔.md", "追踪/时间线.md", "追踪/角色状态.md"]
            for t in tracks:
                if not (r / t).exists():
                    issues.append("missing: " + t)

        chapter_arg = getattr(a, "chapter", None)
        if chapter_arg:
            try:
                ch_num = int(chapter_arg)
                ch_file = None
                for ch in chs:
                    if ("第" + str(ch_num).zfill(2)) in ch.name or ("第" + str(ch_num)) in ch.name:
                        ch_file = ch
                        break
                if ch_file:
                    if not output_json:
                        print("chapter", ch_num, ":", ch_file.name)
                    if getattr(a, "deep", False):
                        body = ch_file.read_text("utf-8")
                        char_count = len(body)
                        if not output_json:
                            print("  chars:", char_count)
                        if char_count < 1500:
                            issues.append("chapter " + str(ch_num) + " too short: " + str(char_count) + " chars (< 1500)")
                        review_file = r / "审查报告" / ("第" + str(ch_num).zfill(2) + "章审查报告.md")
                        if not review_file.exists():
                            issues.append("missing review for chapter " + str(ch_num))
                        commit_file = r / ".story-system" / "commits" / ("chapter_" + str(ch_num).zfill(4) + ".commit.json")
                        if not commit_file.exists():
                            issues.append("missing commit for chapter " + str(ch_num))
                else:
                    issues.append("chapter " + str(ch_num) + " not found in 正文/")
            except ValueError:
                issues.append("invalid chapter number: " + str(chapter_arg))

        if getattr(a, "deep", False) and not chapter_arg:
            ss = r / ".story-system"
            if not ss.exists():
                issues.append("missing: .story-system/ (底本目录)")
            else:
                master = ss / "MASTER_SETTING.json"
                if not master.exists():
                    issues.append("missing: .story-system/MASTER_SETTING.json")

        if output_json:
            print(json.dumps({"issues": issues, "ok": len(issues) == 0, "chapters": len(chs)}, ensure_ascii=False))
        else:
            for i in issues:
                print(i)
            if not issues:
                print("project structure OK")
        return 1 if issues else 0
    except Exception as e:
        message = "diagnose error: " + str(e)
        if output_json:
            print(json.dumps({"issues": [message], "ok": False, "chapters": 0}, ensure_ascii=False))
        else:
            print(message)
            print("tip: check if project path contains special characters")
        return 2

def cmd_contract(a):
    r = find_project(Path(a.project_root))
    if not r:
        print("no project")
        return
    s = load_state(r)
    # load_state 已自动迁移，直接用 'project'
    pi = s.get("project", {})
    d = r / ".story-system"
    for dirname in ("chapters", "volumes", "commits", "reviews", "summaries"):
        (d / dirname).mkdir(parents=True, exist_ok=True)
    current_version = int(s.get("versions", {}).get("baseline_version", 0) or 0)
    m = {
        "schema_version": 1,
        "route": {
            "primary_genre": pi.get("genre", ""),
            "target_platform": pi.get("target_platform", "fanqie"),
        },
        "project": {"book_name": pi.get("book_name", ""), "author": pi.get("author", "")},
    }
    master_file = d / "MASTER_SETTING.json"
    existing = {}
    if master_file.exists():
        try:
            existing = json.loads(master_file.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError):
            existing = {}
    existing_version = int(existing.get("versions", {}).get("baseline_version", 0) or 0)
    same_contract = (
        isinstance(existing, dict)
        and {k: existing.get(k) for k in ("schema_version", "route", "project")} == m
    )
    if same_contract:
        baseline_version = max(1, current_version, existing_version)
    else:
        baseline_version = max(1, current_version + 1, existing_version + 1)
    m["versions"] = {"baseline_version": baseline_version}
    master_file.write_text(json.dumps(m, ensure_ascii=False, indent=2), "utf-8")
    s.setdefault("versions", {})["baseline_version"] = baseline_version
    save_state(r, s)
    print("baseline done:", baseline_version)

def cmd_export(a):
    r = find_project(Path(a.project_root))
    if not r:
        print("no project")
        return
    chs = chapter_files(r)
    if not chs:
        print("no chapters")
        return

    fmt = getattr(a, "format", "txt") or "txt"
    output_arg = getattr(a, "output", None)

    if output_arg:
        out_path = Path(output_arg)
        if not out_path.is_absolute():
            out_path = r / out_path
        if out_path.is_dir() or output_arg.endswith(("/", "\\")):
            suffix = "txt" if fmt in ("txt", "fanqie") else "docx"
            out_path = out_path / (r.name + "_完本." + suffix)
    else:
        out_dir = r / "导出"
        out_dir.mkdir(exist_ok=True)
        suffix = "txt" if fmt in ("txt", "fanqie") else "docx"
        out_path = out_dir / (r.name + "_完本." + suffix)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    parts = [markdown_to_plain_text(ch.read_text("utf-8")) for ch in chs]

    if fmt == "fanqie":
        merged = []
        for ch, body in zip(chs, parts):
            title = ch.stem
            if not title.startswith("第"):
                title = "第" + title + "章"
            merged.append(title + "\n\n" + body + "\n\n\n")
        out_path.write_text("".join(merged), encoding="utf-8")
    elif fmt == "docx":
        try:
            from docx import Document
            doc = Document()
            for ch, body in zip(chs, parts):
                title = ch.stem
                doc.add_heading(title, level=1)
                for para in body.split("\n"):
                    if para.strip():
                        doc.add_paragraph(para)
                doc.add_page_break()
            doc.save(str(out_path))
        except ImportError:
            fallback = out_path.with_suffix(".txt")
            fallback.write_text("\n\n".join(parts), encoding="utf-8")
            print("warning: python-docx not installed, fallback to txt:", fallback)
            out_path = fallback
    else:
        out_path.write_text("\n\n".join(parts), encoding="utf-8")

    print("exported:", out_path)
    print("  chapters:", len(chs))
    print("  total chars:", sum(len(p) for p in parts))

def cmd_commit(a):
    r = find_project(Path(a.project_root))
    if not r:
        print("no project")
        return
    ch = int(a.chapter)
    d = r / ".story-system" / "commits"
    d.mkdir(parents=True, exist_ok=True)
    content_hash = None
    source_path = None
    candidates = [p for p in chapter_files(r) if chapter_number(p) == ch]
    if candidates:
        source_path = candidates[0].relative_to(r).as_posix()
        content_hash = hashlib.sha256(candidates[0].read_bytes()).hexdigest()
    status = "accepted" if source_path else "unverified"
    c = {
        "chapter": ch,
        "timestamp": datetime.now().isoformat(),
        "status": status,
        "source_path": source_path,
        "content_sha256": content_hash,
    }
    artifact_status, data_artifacts, artifact_errors = collect_data_artifacts(
        r, getattr(a, "data_artifacts", None)
    )
    c["data_artifacts_status"] = artifact_status
    c["data_artifacts"] = data_artifacts
    if artifact_errors:
        c["data_artifacts_errors"] = artifact_errors
    (d / f"chapter_{ch:04d}.commit.json").write_text(json.dumps(c, ensure_ascii=False, indent=2), "utf-8")
    st = load_state(r)
    progress = st.setdefault("progress", {})
    if source_path:
        progress["current_chapter"] = max(int(progress.get("current_chapter", 0) or 0), ch)
        progress["last_commit_at"] = c["timestamp"]
    save_state(r, st)
    if not source_path:
        print("warning: chapter not found in 正文/, commit recorded as unverified", file=sys.stderr)

def cmd_review(a):
    r = find_project(Path(a.project_root))
    if not r:
        print("no project")
        return
    rf = Path(a.report_file)
    rf.parent.mkdir(parents=True, exist_ok=True)
    results, results_source, results_error = load_review_results(
        getattr(a, "review_results", None)
    )
    results_verified = review_results_verified(results)
    report_lines = [
        "# 第 " + str(a.chapter) + " 章审查报告",
        "",
        "生成: " + datetime.now().strftime("%Y-%m-%d %H:%M"),
        ""
    ]
    if results_error:
        report_lines.extend([
            "审查结果未加载，指标状态为 unverified。",
            "",
            "原因: " + results_error,
            "",
        ])
        print("warning: " + results_error, file=sys.stderr)
    elif results:
        if isinstance(results, dict):
            for k, v in results.items():
                report_lines.append("## " + str(k))
                report_lines.append("")
                if isinstance(v, str):
                    report_lines.append(v)
                elif isinstance(v, list):
                    for item in v:
                        report_lines.append("- " + str(item))
                else:
                    report_lines.append(str(v))
                report_lines.append("")
        elif isinstance(results, list):
            report_lines.append("## 审查结果")
            report_lines.append("")
            for idx, item in enumerate(results, 1):
                if isinstance(item, dict):
                    report_lines.append("### 条目 " + str(idx))
                    for k, v in item.items():
                        if isinstance(v, list):
                            for sub in v:
                                report_lines.append("- " + str(k) + ": " + str(sub))
                        else:
                            report_lines.append("- " + str(k) + ": " + str(v))
                    report_lines.append("")
                else:
                    report_lines.append(str(idx) + ". " + str(item))
            report_lines.append("")
        else:
            report_lines.append(str(results))
            report_lines.append("")
        if not results_verified:
            report_lines.extend([
                "reviewer 结果声明未完成，指标状态为 unverified。",
                "",
            ])
    else:
        report_lines.append("由 reviewer agent 输出。")
    rf.write_text("\n".join(report_lines), "utf-8")
    print("审查报告:", rf)
    if getattr(a, "save_metrics", False):
        mp = Path(a.metrics_out) if getattr(a, "metrics_out", None) else (
            r / ".novel" / "reviews" / f"chapter_{int(a.chapter):04d}.metrics.json"
        )
        mp.parent.mkdir(parents=True, exist_ok=True)
        metrics = {
            "chapter": int(a.chapter),
            "timestamp": datetime.now().isoformat(),
            "status": "verified" if results_verified else "unverified",
            "results": results if isinstance(results, (dict, list)) else {},
        }
        if results_source:
            metrics["results_source"] = results_source
        if results_error:
            metrics["error"] = results_error
        mp.write_text(json.dumps(metrics, ensure_ascii=False, indent=2), "utf-8")
        print("metrics saved:", mp)


def cmd_projections(a):
    r = find_project(Path(a.project_root))
    if not r:
        print("no project")
        return 1
    try:
        chapter = int(a.chapter)
    except (TypeError, ValueError):
        print("invalid chapter number: " + str(a.chapter), file=sys.stderr)
        return 1

    commit_file = r / ".story-system" / "commits" / f"chapter_{chapter:04d}.commit.json"
    if not commit_file.exists():
        print("commit not found: " + str(commit_file), file=sys.stderr)
        return 1
    try:
        commit = json.loads(commit_file.read_text(encoding="utf-8"))
        action = getattr(a, "action", "status")
        projection = commit.get("projection_status", {})
        if action == "status":
            payload = {
                "chapter": chapter,
                "projection_status": projection,
                "summary_file": str(r / ".story-system" / "summaries" / f"ch{chapter:04d}.md"),
            }
            print(json.dumps(payload, ensure_ascii=False))
            return 0 if projection.get("summary") == "done" else 1

        summary_file = project_chapter_summary(r, chapter, commit)
        projection = {
            "state": "done",
            "summary": "done",
            "index": "skipped",
            "memory": "skipped",
            "vector": "skipped",
        }
        commit["projection_status"] = projection
        commit["projection_updated_at"] = datetime.now().isoformat()
        commit_file.write_text(json.dumps(commit, ensure_ascii=False, indent=2), "utf-8")
        log_file = r / ".story-system" / "projection_log.jsonl"
        log_file.parent.mkdir(parents=True, exist_ok=True)
        with log_file.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps({
                "timestamp": datetime.now().isoformat(),
                "chapter": chapter,
                "action": action,
                "content_sha256": commit.get("content_sha256"),
                "projection_status": projection,
            }, ensure_ascii=False) + "\n")
        print("projection refreshed:", summary_file)
        print("  model-managed projections: index=skipped, memory=skipped, vector=skipped")
        return 0
    except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
        print("projection error: " + str(exc), file=sys.stderr)
        return 2

def cmd_memory(a):
    r = find_project(Path(a.project_root))
    if not r:
        print("no project")
        return
    mf = r / ".novel" / "project_memory.json"
    if not mf.exists():
        mf.write_text(json.dumps({"patterns": []}, ensure_ascii=False, indent=2), "utf-8")
    try:
        mem = json.loads(mf.read_text("utf-8"))
    except Exception:
        mem = {"patterns": []}
    patterns = mem.setdefault("patterns", [])

    if a.action == "add-pattern":
        p = {"pattern_type": a.pattern_type, "description": a.description, "importance": a.importance or "medium"}
        # 去重：pattern_type + description 完全相同则跳过
        for existing in patterns:
            if existing.get("pattern_type") == p["pattern_type"] and existing.get("description") == p["description"]:
                print("pattern already exists, skipped")
                return
        patterns.append(p)
        mf.write_text(json.dumps(mem, ensure_ascii=False, indent=2), "utf-8")
        print("pattern saved: " + str(p["pattern_type"]))

    elif a.action == "list-patterns":
        if not patterns:
            print("(no patterns yet)")
            return
        for idx, p in enumerate(patterns, 1):
            print(f"{idx}. [{p.get('pattern_type', '?')}] {p.get('description', '?')[:60]} (importance: {p.get('importance', 'medium')})")

    elif a.action == "search-pattern":
        if not patterns:
            print("(no patterns yet)")
            return
        keyword = (a.keyword or "").lower()
        ptype = (a.pattern_type or "").lower()
        hits = []
        for idx, p in enumerate(patterns, 1):
            desc = p.get("description", "").lower()
            pt = p.get("pattern_type", "").lower()
            if keyword and keyword in desc:
                hits.append((idx, p))
            elif ptype and ptype == pt:
                hits.append((idx, p))
            elif not keyword and not ptype:
                hits.append((idx, p))
        if not hits:
            print("(no matches)")
            return
        for idx, p in hits:
            print(f"{idx}. [{p.get('pattern_type', '?')}] {p.get('description', '?')} (importance: {p.get('importance', 'medium')})")

    elif a.action == "delete-pattern":
        if not patterns:
            print("(no patterns yet)")
            return
        target = a.target
        # 优先按索引删除
        try:
            idx = int(target) - 1
            if 0 <= idx < len(patterns):
                removed = patterns.pop(idx)
                mf.write_text(json.dumps(mem, ensure_ascii=False, indent=2), "utf-8")
                print(f"deleted: [{removed.get('pattern_type', '?')}] {removed.get('description', '?')[:50]}")
                return
        except ValueError:
            pass
        # 按类型删除（删除该类型所有 pattern）
        before = len(patterns)
        patterns[:] = [p for p in patterns if p.get("pattern_type") != target]
        after = len(patterns)
        if before != after:
            mf.write_text(json.dumps(mem, ensure_ascii=False, indent=2), "utf-8")
            print(f"deleted {before - after} pattern(s) of type: {target}")
        else:
            print(f"no pattern found with type: {target}")

def main():
    p = argparse.ArgumentParser()
    p.add_argument("--project-root", default=os.getcwd())
    s = p.add_subparsers(dest="cmd")
    pi = s.add_parser("init", help="create a new project")
    pi.add_argument("--title", default="")
    pi.add_argument("--author", default="")
    pi.add_argument("--genre", default="")
    pi.set_defaults(func=cmd_init)
    s.add_parser("where").set_defaults(func=cmd_where)
    ps = s.add_parser("project-status")
    ps.add_argument("--format", choices=["summary", "json"], default="summary")
    ps.set_defaults(func=cmd_status)
    pd = s.add_parser("doctor")
    pd.add_argument("--format", choices=["text", "json"], default="text")
    pd.add_argument("--chapter", default=None, help="检查指定章节")
    pd.add_argument("--deep", action="store_true", help="深度检查（含底本一致性）")
    pd.set_defaults(func=cmd_doctor)
    s.add_parser("init-contract").set_defaults(func=cmd_contract)
    pe = s.add_parser("export")
    pe.add_argument("--format", choices=["txt", "docx", "fanqie"], default="txt",
                    help="导出格式：txt 纯文本 / docx Word / fanqie 番茄平台规范")
    pe.add_argument("--output", default=None,
                    help="输出路径（目录或文件名）；不指定则写到 导出/ 下")
    pe.set_defaults(func=cmd_export)
    p6 = s.add_parser("chapter-commit")
    p6.add_argument("--chapter", required=True)
    p6.add_argument(
        "--data-artifacts",
        default=None,
        help="data-agent JSON 文件或目录；默认读取 .novel/tmp/ 下的三份 artifact",
    )
    p6.set_defaults(func=cmd_commit)
    p7 = s.add_parser("review-pipeline")
    p7.add_argument("--chapter", required=True)
    p7.add_argument("--report-file", required=True)
    p7.add_argument("--review-results")
    p7.add_argument("--metrics-out")
    p7.add_argument("--save-metrics", action="store_true")
    p7.set_defaults(func=cmd_review)
    p9 = s.add_parser("projections")
    p9.add_argument("action", choices=["status", "retry", "replay"], default="status")
    p9.add_argument("--chapter", required=True)
    p9.set_defaults(func=cmd_projections)
    p8 = s.add_parser("project-memory")
    p8.add_argument("action", choices=["add-pattern", "list-patterns", "search-pattern", "delete-pattern"],
                    help="add-pattern / list-patterns / search-pattern / delete-pattern")
    p8.add_argument("--pattern-type", default=None, help="模式类型（add 必填，search/delete 可选）")
    p8.add_argument("--description", default=None, help="模式描述（add 必填）")
    p8.add_argument("--importance", default="medium")
    p8.add_argument("--keyword", default=None, help="搜索关键词（search-pattern 用）")
    p8.add_argument("--target", default=None, help="删除目标：索引数字或类型名（delete-pattern 用）")
    p8.set_defaults(func=cmd_memory)
    a = p.parse_args()
    if hasattr(a, "func"):
        result = a.func(a)
        if isinstance(result, int):
            sys.exit(result)
    else:
        p.print_help()

if __name__ == "__main__":
    main()
